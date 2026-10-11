"""A kept cover letter as a file: PDF, DOCX or DOC. Named like the CV (Company_Role_Name), plus
_Cover_Letter so the two don't overwrite each other when both are saved to one folder."""
import datetime as dt
import html
import re
import subprocess
import tempfile
from pathlib import Path

from .render import in_thread, print_html

FORMATS = {"pdf": "application/pdf",
           "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
           "doc": "application/msword"}


def file_stem(profile: dict, company: str, role: str) -> str:
    from .tailor import pdf_name
    return pdf_name(profile, company, role)[:-len(".pdf")] + "_Cover_Letter"


def _contact(basics: dict) -> list[str]:
    return [x for x in (basics.get("email"), basics.get("phone"), basics.get("location"), basics.get("linkedin")) if x]


def _paragraphs(text: str) -> list[list[str]]:
    """Blank-line separated paragraphs; single newlines inside one (the sign-off) stay line breaks."""
    return [[ln.strip() for ln in p.strip().splitlines()] for p in re.split(r"\n\s*\n", text.strip()) if p.strip()]


def _date() -> str:
    d = dt.date.today()
    return f"{d.day} {d.strftime('%B %Y')}"


def to_docx(text: str, basics: dict, out: Path) -> None:
    from docx import Document
    from docx.shared import Cm, Pt
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(0), 1.15
    p = d.add_paragraph(); r = p.add_run(basics.get("name", "")); r.bold, r.font.size = True, Pt(16)
    if _contact(basics):
        d.add_paragraph("  |  ".join(_contact(basics)))
    d.add_paragraph().paragraph_format.space_after = Pt(6)
    d.add_paragraph(_date()).paragraph_format.space_after = Pt(12)
    for lines in _paragraphs(text):
        p = d.add_paragraph()
        for i, ln in enumerate(lines):
            if i:
                p.add_run().add_break()
            p.add_run(ln)
        p.paragraph_format.space_after = Pt(10)
    d.core_properties.author, d.core_properties.title = basics.get("name", ""), "Cover letter"
    d.save(str(out))


def _rtf(text: str, basics: dict) -> str:
    def esc(s: str) -> str:
        bs = chr(92)
        return "".join(bs + c if c in "{}" + bs else c if ord(c) < 128 else f"{bs}u{ord(c)}?" for c in s)
    out = [r"{\rtf1\ansi\deff0{\fonttbl{\f0 Calibri;}}\paperw11906\paperh16838\margl1417\margr1417\margt1304\margb1304\f0\fs22"]
    out.append(r"{\b\fs32 " + esc(basics.get("name", "")) + r"}\par")
    if _contact(basics):
        out.append(esc("  |  ".join(_contact(basics))) + r"\par")
    out.append(r"\par " + esc(_date()) + r"\par\par")
    for lines in _paragraphs(text):
        out.append(r"\line ".join(esc(ln) for ln in lines) + r"\par\par")
    return "\n".join(out) + "}"


def _soffice() -> str | None:
    import shutil
    for c in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if c and Path(c).exists():
            return c
    return None


def to_doc(text: str, basics: dict, out: Path) -> None:
    """A real Word 97-2003 file when LibreOffice is installed; otherwise RTF, which every version of
    Word opens and which keeps the .doc name an upload form asks for. (Word itself isn't driven: an
    unlicensed copy stops on activation dialogs.)"""
    exe = _soffice()
    if exe:
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
            src = Path(tmp) / "letter.docx"
            to_docx(text, basics, src)
            try:
                subprocess.run([exe, "--headless", "--convert-to", "doc:MS Word 97", "--outdir", tmp, str(src)],
                               check=True, capture_output=True, timeout=90)
                made = Path(tmp) / "letter.doc"
                if made.exists():
                    out.write_bytes(made.read_bytes())
                    return
            except (OSError, subprocess.SubprocessError):
                pass
    out.write_text(_rtf(text, basics), "ascii")


def to_pdf(text: str, basics: dict, out: Path) -> None:
    paras = "".join("<p>" + "<br>".join(html.escape(ln) for ln in lines) + "</p>" for lines in _paragraphs(text))
    contact = html.escape("  |  ".join(_contact(basics)))
    page = f"""<!doctype html><html lang="en-GB"><head><meta charset="utf-8"><title>Cover letter - {html.escape(basics.get('name', ''))}</title>
<style>@page {{ size: A4; margin: 23mm 25mm; }} * {{ margin: 0; padding: 0; }}
html {{ font: 11pt/1.4 Calibri, Carlito, Arial, sans-serif; color: #000; font-variant-ligatures: none; }}
h1 {{ font-size: 16pt; }} .c {{ margin-bottom: 18pt; }} .d {{ margin-bottom: 14pt; }} p {{ margin-bottom: 10pt; }}</style></head>
<body><h1>{html.escape(basics.get('name', ''))}</h1><div class="c">{contact}</div><div class="d">{_date()}</div>{paras}</body></html>"""
    print_html(page, out)
    from .render import set_metadata
    set_metadata(out, "Cover letter", basics.get("name", ""), "Cover letter")


def build(fmt: str, text: str, basics: dict, out: Path) -> None:
    if fmt == "docx":
        to_docx(text, basics, out)
    elif fmt == "doc":
        to_doc(text, basics, out)
    elif fmt == "pdf":
        in_thread(to_pdf, text, basics, out)
    else:
        raise ValueError(fmt)
